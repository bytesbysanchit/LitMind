"""Backend adapter for LitMind's annotated Dattani Play mode.

This module produces the same *kind* of analysis payload as novel analysis
(character summary, contexts, weighted edges, section-level metrics), while
keeping play-specific semantics explicit:
- character counts are annotated speaker turns, not NER mention counts;
- interaction edges are adjacent speaker-turn transitions, not relationships;
- semantic relations are emitted only when explicitly documented in the corpus.

It deliberately does not import or modify src/nlp_processor.py or src/prev.py.
"""
from __future__ import annotations

from collections import Counter, defaultdict
from pathlib import Path
import re
from typing import Any

import pandas as pd
from openpyxl import load_workbook

PLAY_NAMES = {"D01": "Tara", "D02": "Dance Like a Man"}


def _clean(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def _norm(value: Any) -> str:
    return re.sub(r"[^A-Z0-9]+", "", _clean(value).upper())


def _headers(ws) -> dict[str, int]:
    row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    return {_clean(value): idx for idx, value in enumerate(row) if value is not None}


def _optional_sentiment(text: str) -> tuple[float | None, str]:
    """Return VADER compound score and label, or unavailable if not installed."""
    try:
        from vaderSentiment.vaderSentiment import SentimentIntensityAnalyzer
        analyzer = SentimentIntensityAnalyzer()
        score = float(analyzer.polarity_scores(text)["compound"])
        label = "Positive" if score >= 0.05 else "Negative" if score <= -0.05 else "Neutral"
        return score, label
    except Exception:
        return None, "Unavailable"


def _documented_dance_relations(roster: dict[str, dict[str, Any]]) -> list[dict[str, Any]]:
    """Extract only explicit pairwise relations grounded in CHARACTER_MASTER Notes.

    The five rules below are tied to the corpus's documented character notes.
    They are not inferred from adjacency or NER and are not claimed exhaustive.
    """
    names = {name.casefold(): name for name in roster}
    wanted = [
        ("Jairaj", "Ratna", "SPOUSE", "Ratna's CHARACTER_MASTER Notes identify her as Jairaj's wife."),
        ("Lata", "Viswas", "FIANCE", "Viswas's CHARACTER_MASTER Notes identify him as Lata's fiancé."),
        ("Jairaj", "Lata", "PARENT_CHILD", "Lata's CHARACTER_MASTER Notes identify her as Jairaj's daughter."),
        ("Ratna", "Lata", "PARENT_CHILD", "Lata's CHARACTER_MASTER Notes identify her as Ratna's daughter."),
        ("Amritlal", "Jairaj", "PARENT_CHILD", "Amritlal's CHARACTER_MASTER Notes identify him as Jairaj's father."),
    ]
    output = []
    for a, b, relation, evidence in wanted:
        if a.casefold() in names and b.casefold() in names:
            output.append({
                "character_a": names[a.casefold()],
                "character_b": names[b.casefold()],
                "relation_label": relation,
                "evidence_source": "CHARACTER_MASTER.Notes",
                "evidence_text": evidence,
                "evidence_status": "DOCUMENTED_IN_CORPUS_NOTES",
                "is_exhaustive_gold": False,
            })
    return output


def analyze_play(corpus_path: str | Path, play_id: str) -> dict[str, Any]:
    """Analyze one annotated play and return a merge-ready LitMind payload.

    Args:
        corpus_path: path to Dattani_Tara_Corpus_Final.xlsx.
        play_id: D01 (Tara) or D02 (Dance Like a Man).

    The result intentionally separates interaction graph edges from semantic
    relationship evidence so a downstream UI cannot silently conflate them.
    """
    play_id = _clean(play_id).upper()
    if play_id not in PLAY_NAMES:
        raise ValueError(f"Unsupported play_id {play_id!r}; expected D01 or D02")

    wb = load_workbook(Path(corpus_path), read_only=True, data_only=True)
    try:
        required = {"CHARACTER_MASTER", "UTTERANCE_CORPUS"}
        if not required.issubset(wb.sheetnames):
            raise ValueError(f"Corpus workbook must contain sheets: {sorted(required)}")
        master_ws = wb["CHARACTER_MASTER"]
        utter_ws = wb["UTTERANCE_CORPUS"]
        mh, uh = _headers(master_ws), _headers(utter_ws)
        needed_m = {"Character_ID", "Play_ID", "Character_Name", "Source_Speaker_Label", "Notes"}
        needed_u = {"Speech_ID", "Play_ID", "Play", "Act", "Speaker", "Speech_OCR", "QA_Flag"}
        if not needed_m.issubset(mh):
            raise ValueError(f"CHARACTER_MASTER missing columns: {sorted(needed_m - set(mh))}")
        if not needed_u.issubset(uh):
            raise ValueError(f"UTTERANCE_CORPUS missing columns: {sorted(needed_u - set(uh))}")

        roster: dict[str, dict[str, Any]] = {}
        speaker_lookup: dict[str, str] = {}
        notes_lookup: dict[str, str] = {}
        for row in master_ws.iter_rows(min_row=2, values_only=True):
            if _clean(row[mh["Play_ID"]]).upper() != play_id:
                continue
            name = _clean(row[mh["Character_Name"]])
            label = _clean(row[mh["Source_Speaker_Label"]])
            if not name or not label:
                continue
            roster[name] = {
                "character_id": _clean(row[mh["Character_ID"]]),
                "name": name,
                "speaker_label": label,
                "speaker_turns": 0,
                "first_context": "",
                "qa_flags": Counter(),
            }
            speaker_lookup[_norm(label)] = name
            notes_lookup[name] = _clean(row[mh["Notes"]])
        if not roster:
            raise ValueError(f"No reference characters found for {play_id}")

        # Novel-like output fields, but their play-specific meaning is explicitly tagged.
        character_counts: Counter[str] = Counter()
        character_context: dict[str, str] = {}
        transition_direction_counts: Counter[tuple[str, str, str]] = Counter()
        transition_examples: defaultdict[tuple[str, str], list[str]] = defaultdict(list)
        sections: dict[str, dict[str, Any]] = {}
        unmatched_speakers: Counter[str] = Counter()
        blank_speakers = 0
        utterance_rows = 0
        unique_speech_ids: set[str] = set()
        previous_by_act: dict[str, dict[str, str]] = {}
        mention_candidates: list[dict[str, Any]] = []

        # Read rows in source order: consecutive speakers define interaction transitions.
        for row in utter_ws.iter_rows(min_row=2, values_only=True):
            if _clean(row[uh["Play_ID"]]).upper() != play_id:
                continue
            utterance_rows += 1
            speech_id = _clean(row[uh["Speech_ID"]])
            if speech_id:
                unique_speech_ids.add(speech_id)
            act = _clean(row[uh["Act"]]) or "Act not specified"
            speaker_label = _clean(row[uh["Speaker"]])
            text = _clean(row[uh["Speech_OCR"]])
            qa_flag = _clean(row[uh["QA_Flag"]]) or "No QA flag"
            sec = sections.setdefault(act, {"Act": act, "Utterance rows": 0, "Words": 0, "Text parts": [], "Sentiment compound": None, "Sentiment label": "Unavailable"})
            sec["Utterance rows"] += 1
            sec["Words"] += len(re.findall(r"\b\w+\b", text))
            if text:
                sec["Text parts"].append(text)

            if not speaker_label:
                blank_speakers += 1
                continue
            canonical = speaker_lookup.get(_norm(speaker_label))
            if canonical is None:
                unmatched_speakers[speaker_label] += 1
                continue

            character_counts[canonical] += 1
            roster[canonical]["speaker_turns"] += 1
            roster[canonical]["qa_flags"][qa_flag] += 1
            if not roster[canonical]["first_context"] and text:
                roster[canonical]["first_context"] = text[:240]
            character_context.setdefault(canonical, text[:240])

            # Alias candidates are generated separately from authoritative speaker labels.
            # They are not auto-accepted as mentions or nodes.
            for name in roster:
                pattern = re.compile(rf"(?<!\w){re.escape(name)}(?!\w)", re.IGNORECASE)
                for match in pattern.finditer(text):
                    mention_candidates.append({
                        "Speech_ID": speech_id,
                        "Act": act,
                        "Speaker": speaker_label,
                        "Candidate_character": name,
                        "Matched_surface": match.group(0),
                        "Context": text[max(0, match.start()-100):min(len(text), match.end()+140)],
                        "QA_Flag": qa_flag,
                        "Evidence_status": "ALIAS_CANDIDATE_REVIEW_REQUIRED",
                    })

            prev = previous_by_act.get(act)
            current = {"speaker": canonical, "speech_id": speech_id}
            if prev and prev["speaker"] != canonical:
                a, b = sorted((prev["speaker"], canonical), key=str.casefold)
                direction = "A_then_B" if (prev["speaker"], canonical) == (a, b) else "B_then_A"
                transition_direction_counts[(a, b, direction)] += 1
                if len(transition_examples[(a, b)]) < 5:
                    transition_examples[(a, b)].append(f'{prev["speech_id"]} → {speech_id}')
            previous_by_act[act] = current

        interaction_edges = Counter()
        for (a, b, _direction), count in transition_direction_counts.items():
            interaction_edges[(a, b)] += count
        interaction_edge_rows = []
        for (a, b), count in sorted(interaction_edges.items(), key=lambda item: (-item[1], item[0][0].casefold(), item[0][1].casefold())):
            interaction_edge_rows.append({
                "character_a": a,
                "character_b": b,
                "transition_count": count,
                "a_then_b": transition_direction_counts[(a, b, "A_then_B")],
                "b_then_a": transition_direction_counts[(a, b, "B_then_A")],
                "examples": transition_examples[(a, b)],
                "edge_type": "ADJACENT_SPEAKER_TURN_INTERACTION",
                "semantic_relationship": False,
            })

        # Per-Act sentiment is optional and uses the same VADER family as novel mode.
        sentiment_available = False
        section_rows = []
        for act, sec in sections.items():
            body = "\n".join(sec.pop("Text parts"))
            score, label = _optional_sentiment(body) if body else (None, "Unavailable")
            if score is not None:
                sentiment_available = True
            section_rows.append({"Act": act, "Utterance rows": sec["Utterance rows"], "Words": sec["Words"], "Sentiment compound": score, "Sentiment label": label})
        section_frame = pd.DataFrame(section_rows)

        # Relationship evidence remains a separate, provenance-labelled channel.
        semantic_edges = _documented_dance_relations(roster) if play_id == "D02" else []
        character_rows = []
        for name, item in sorted(roster.items(), key=lambda kv: (-kv[1]["speaker_turns"], kv[0].casefold())):
            character_rows.append({
                "character_id": item["character_id"], "name": name,
                "speaker_label": item["speaker_label"],
                "speaker_turns": item["speaker_turns"],
                "first_context": item["first_context"],
                "qa_flags": dict(item["qa_flags"]),
                "selection_status": "SELECTED_REFERENCE_CHARACTER" if item["speaker_turns"] else "REVIEW_ZERO_SPEAKER_TURNS",
                "character_master_notes": notes_lookup.get(name, ""),
            })

        return {
            "schema_version": "1.0",
            "mode": "play",
            "play_id": play_id,
            "title": PLAY_NAMES[play_id],
            "summary": {
                "utterance_rows": utterance_rows,
                "unique_speech_ids": len(unique_speech_ids),
                "reference_characters": len(roster),
                "characters_with_speaker_turns": sum(c["speaker_turns"] > 0 for c in character_rows),
                "mapped_speaker_turns": sum(character_counts.values()),
                "blank_speaker_rows": blank_speakers,
                "unmapped_speaker_labels": dict(unmatched_speakers),
                "status": "PASS" if not unmatched_speakers and not blank_speakers and sum(character_counts.values()) == utterance_rows else "REVIEW",
            },
            # Shared display concepts with novel mode, with an explicit count unit.
            "character_counts": character_counts,
            "character_context": character_context,
            "character_count_unit": "annotated_speaker_turns",
            "characters": character_rows,
            "interaction_edges": interaction_edges,
            "interaction_edge_rows": interaction_edge_rows,
            "interaction_edge_semantics": "Adjacent annotated speaker turns only; not semantic relationships.",
            "semantic_relation_edges": semantic_edges,
            "semantic_relation_status": "DOCUMENTED_CORPUS_NOTES_ONLY" if semantic_edges else "NO_PAIRWISE_RELATION_EDGES_EMITTED",
            "mention_candidates": mention_candidates,
            "mention_candidate_status": "REVIEW_REQUIRED_NOT_GOLD_MENTIONS",
            "sections": section_frame,
            "sentiment_available": sentiment_available,
            "unmapped_speaker_labels": dict(unmatched_speakers),
            "notes": [
                "Speaker turn counts are not NER mention counts.",
                "Interaction edges are adjacent-turn transitions, not semantic relationships.",
                "Semantic relations are emitted only from explicit corpus notes; the list is not exhaustive gold.",
                "Alias mention candidates require review and may include metadata or stage-direction contexts.",
            ],
        }
    finally:
        wb.close()


def analyze_all_plays(corpus_path: str | Path) -> dict[str, Any]:
    """Analyze both annotated plays with a single backend entry point."""
    results = [analyze_play(corpus_path, pid) for pid in ("D01", "D02")]
    return {
        "schema_version": "1.0",
        "mode": "play_collection",
        "results": results,
        "summary": {
            "plays": len(results),
            "utterance_rows": sum(r["summary"]["utterance_rows"] for r in results),
            "reference_characters": sum(r["summary"]["reference_characters"] for r in results),
            "interaction_edges": sum(len(r["interaction_edge_rows"]) for r in results),
            "status": "PASS" if all(r["summary"]["status"] == "PASS" for r in results) else "REVIEW",
        },
    }
