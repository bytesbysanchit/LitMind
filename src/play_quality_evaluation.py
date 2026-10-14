"""LitMind Play Corpus: automated QA + evaluation-readiness report.
Does not import or modify frozen nlp_processor.py or prev.py.
Does not create independent gold labels or claim model accuracy.

Run from D:\\LitMind:
  python src/play_quality_evaluation.py --corpus Dattani_Tara_Corpus_Final.xlsx --out reports/play_quality_evaluation.xlsx
Requires: openpyxl
"""
from pathlib import Path
from collections import Counter, defaultdict
import argparse, re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

PLAY_NAMES={'D01':'Tara','D02':'Dance Like a Man'}
def clean(v): return re.sub(r'\\s+',' ',str(v or '')).strip()
def norm(v): return clean(v).casefold()
def play_key(pid,name=''):
    p=clean(pid).upper(); t=(p+' '+clean(name)).casefold()
    if p=='D01' or 'tara' in t: return 'Tara'
    if p=='D02' or 'dance like a man' in t or ('dance' in t and 'man' in t): return 'Dance Like a Man'
    return None

def main():
    ap=argparse.ArgumentParser()
    ap.add_argument('--corpus',required=True)
    ap.add_argument('--out',default='reports/play_quality_evaluation.xlsx')
    args=ap.parse_args(); cp=Path(args.corpus); op=Path(args.out); op.parent.mkdir(parents=True,exist_ok=True)
    if not cp.exists(): raise SystemExit(f'Corpus not found: {cp}')
    src=openpyxl.load_workbook(cp,read_only=True,data_only=True)
    required={'CHARACTER_MASTER','UTTERANCE_CORPUS','CORPUS_SUMMARY','SOURCE_QA'}
    missing=required-set(src.sheetnames)
    if missing: raise SystemExit('Missing required sheets: '+', '.join(sorted(missing)))
    def rows(sheet):
        data=list(src[sheet].iter_rows(values_only=True)); return [clean(x) for x in data[0]], data[1:]
    mh,mrows=rows('CHARACTER_MASTER'); uh,urows=rows('UTTERANCE_CORPUS'); sh,srows=rows('CORPUS_SUMMARY'); qh,qrows=rows('SOURCE_QA')
    mi={x:i for i,x in enumerate(mh)}; ui={x:i for i,x in enumerate(uh)}; si={x:i for i,x in enumerate(sh)}
    for c in ['Play_ID','Character_Name','Source_Speaker_Label']:
        if c not in mi: raise SystemExit(f'CHARACTER_MASTER missing {c}')
    for c in ['Speech_ID','Play_ID','Play','Act','Speaker','Speech_OCR','QA_Flag']:
        if c not in ui: raise SystemExit(f'UTTERANCE_CORPUS missing {c}')
    refs=defaultdict(set); labels=defaultdict(set); notes={}
    for r in mrows:
        if not r: continue
        p=play_key(r[mi['Play_ID']]); name=clean(r[mi['Character_Name']]); lab=clean(r[mi['Source_Speaker_Label']]).upper()
        if p and name:
            refs[p].add(name)
            if lab: labels[p].add(lab)
            notes[(p,name)]=clean(r[mi['Notes']]) if 'Notes' in mi else ''
    data=[]
    for r in urows:
        if not r: continue
        p=play_key(r[ui['Play_ID']],r[ui['Play']])
        if not p: continue
        d={k:(r[i] if i<len(r) else None) for k,i in ui.items()}
        d['play_key']=p; d['speaker_norm']=clean(d['Speaker']).upper(); d['text']=clean(d['Speech_OCR']); d['speech_id']=clean(d['Speech_ID']); d['qa']=clean(d['QA_Flag'])
        data.append(d)
    wb=openpyxl.Workbook(); summary=wb.active; summary.title='Summary'
    summary.append(['Check','Result','Interpretation'])
    summary.append(['Source corpus',str(cp),'Original source; this script does not modify it.'])
    summary.append(['Utterance rows analyzed',len(data),'Rows from UTTERANCE_CORPUS for the two target plays.'])
    # count by play
    playcounts=Counter(d['play_key'] for d in data)
    for p in ['Tara','Dance Like a Man']:
        summary.append([f'{p}: reference characters',len(refs[p]),'From CHARACTER_MASTER; not model accuracy.'])
        summary.append([f'{p}: utterance rows',playcounts[p],'Count from UTTERANCE_CORPUS.'])
    ids=Counter(d['speech_id'] for d in data if d['speech_id'])
    duplicates={k:v for k,v in ids.items() if v>1}
    blankspeaker=[d for d in data if not d['speaker_norm']]
    blanktext=[d for d in data if not d['text']]
    unknown=[d for d in data if d['speaker_norm'] and d['speaker_norm'] not in labels[d['play_key']]]
    summary.append(['Duplicate nonblank Speech_IDs',len(duplicates),'Investigate; does not automatically mean data error.'])
    summary.append(['Blank speaker labels',len(blankspeaker),'Needs review before speaker evaluation.'])
    summary.append(['Blank utterance text',len(blanktext),'Needs review before text-based evaluation.'])
    summary.append(['Speaker labels not found in CHARACTER_MASTER',len(unknown),'Review mapping; not automatically incorrect.'])
    summary.append(['Independent gold labels present in source?', 'NO','Source labels are references/annotations, not independently adjudicated evaluation labels.'])
    summary.append(['Accuracy / precision / recall / F1 calculated?', 'NO','Metrics are intentionally withheld until independent gold labels and predictions are available.'])
    # row-level QA
    qa=wb.create_sheet('Utterance_QA_Flags'); qa.append(['Play','Speech_ID','Act','Speaker label','QA_Flag','Issue flags','Text preview','Recommended action'])
    for d in data:
        issues=[]
        if not d['speaker_norm']: issues.append('BLANK_SPEAKER')
        elif d['speaker_norm'] not in labels[d['play_key']]: issues.append('UNMAPPED_SPEAKER_LABEL')
        if not d['text']: issues.append('BLANK_TEXT')
        if d['speech_id'] and ids[d['speech_id']]>1: issues.append('DUPLICATE_SPEECH_ID')
        text=d['text']; low=text.casefold()
        if any(x in low for x in ['copyright','all rights reserved','a drama in two acts by','title of the play']): issues.append('POSSIBLE_FRONTMATTER_IN_DIALOGUE')
        if re.search(r'\\([^)]{1,100}\\)',text) and re.search(r'\\b(pauses?|moves|looks|sits|stands|laughs|smiles|cries|music fades|picks up|puts down)\\b',text,re.I): issues.append('STAGE_DIRECTION_MIXED_IN_TEXT_POSSIBLE')
        if d['qa']: issues.append('SOURCE_QA_FLAG_PRESENT')
        if issues:
            qa.append([d['play_key'],d['speech_id'],clean(d['Act']),d['speaker_norm'],d['qa'],'; '.join(issues),text[:350],'Review source/context; do not auto-correct from this flag alone.'])
    # QA flag distribution
    qdist=wb.create_sheet('QA_Flag_Distribution'); qdist.append(['Play','QA_Flag','Rows'])
    dist=Counter((d['play_key'],d['qa'] or '(blank)') for d in data)
    for (p,q),n in sorted(dist.items()): qdist.append([p,q,n])
    # speaker counts
    sw=wb.create_sheet('Speaker_Label_Counts'); sw.append(['Play','Speaker label','Rows','Known in CHARACTER_MASTER?'])
    sc=Counter((d['play_key'],d['speaker_norm'] or '(blank)') for d in data)
    for (p,l),n in sorted(sc.items()): sw.append([p,l,n,'YES' if l in labels[p] else 'NO'])
    # annotations conflict report from Tara sheet if present
    rel=wb.create_sheet('Tara_Relation_Audit'); rel.append(['Speaker','Target','Relation label','Rows','Speech_ID examples','Audit flag','Interpretation'])
    if 'TARA_FULL_ANNOTATION_V4' in src.sheetnames:
        th, tr=rows('TARA_FULL_ANNOTATION_V4'); ti={x:i for i,x in enumerate(th)}
        req=['Speaker','Speech_ID','Masculinity_Target','Secondary_Target','Target_Relation','Secondary_Target_Relation','Primary_Text','Source_QA_Flag']
        grouped=defaultdict(list)
        for r in tr:
            if not r: continue
            def val(k): return clean(r[ti[k]]) if k in ti and ti[k]<len(r) else ''
            speaker=val('Speaker')
            for targetcol,relcol in [('Masculinity_Target','Target_Relation'),('Secondary_Target','Secondary_Target_Relation')]:
                target=val(targetcol); label=val(relcol).upper()
                if target and target not in {'ND','SELF','—','-'} and label and label not in {'ND','—','-'}:
                    # preserve source target string; no assumption that target is a character
                    grouped[(speaker,target,label)].append((val('Speech_ID'),val('Source_QA_Flag'),val('Primary_Text')))
        pairlabels=defaultdict(set)
        for speaker,target,label in grouped: pairlabels[(speaker,target)].add(label)
        for (speaker,target,label),items in sorted(grouped.items()):
            flags=[]
            if label in {'RELATIONAL','OTHER_MALE','OTHER_FEMALE','OTHER'}: flags.append('GENERIC_LABEL_REVIEW')
            if len(pairlabels[(speaker,target)])>1: flags.append('MULTIPLE_LABELS_FOR_SAME_SPEAKER_TARGET')
            if any(x[1] for x in items): flags.append('SOURCE_QA_FLAG_PRESENT_IN_EVIDENCE')
            rel.append([speaker,target,label,len(items),'','; '.join(flags) if flags else 'NO_AUTOMATED_FLAG','Source annotation only; not independent gold truth.'])
            # correct examples cell after row appended
            rel.cell(rel.max_row,5).value='; '.join(list(dict.fromkeys(x[0] for x in items if x[0]))[:8])
    # relation QA counts summary
    relsummary=wb.create_sheet('Relation_Audit_Summary'); relsummary.append(['Metric','Count','Meaning'])
    relsummary.append(['Distinct speaker-target-label groups',max(0,rel.max_row-1),'Existing Tara annotations grouped for inspection.'])
    relsummary.append(['Groups flagged for review',sum(1 for row in list(rel.values)[1:] if row[5] not in ('NO_AUTOMATED_FLAG',None)),'Automated flags are review prompts, not proof of annotation error.'])
    # source QA findings
    sourceqa=wb.create_sheet('Source_QA_Findings'); sourceqa.append(['QA_Item','Finding'])
    for r in qrows:
        if r: sourceqa.append([clean(r[0]) if len(r)>0 else '',clean(r[1]) if len(r)>1 else ''])
    # Evaluation protocol
    protocol=wb.create_sheet('Evaluation_Readiness'); protocol.append(['Component','Can report now?','Reason / next requirement'])
    for row in [
      ('Corpus coverage / counts','YES, descriptively','Counts describe source workbook; reconcile any CORPUS_SUMMARY discrepancy explicitly.'),
      ('Speaker label coverage','YES, as corpus consistency check','Comparing speaker labels with CHARACTER_MASTER checks mapping only; it is not model speaker-identification accuracy.'),
      ('NER precision/recall/F1','NO','Need independently labeled person-mention spans on a held-out sample and saved model predictions.'),
      ('Character alias resolution metrics','NO','Need independent mention-to-character gold mapping, including ambiguous/common-name cases.'),
      ('Relationship precision/recall/F1','NO','Need independent pairwise relation labels, a defined relation schema, and model predictions on the same held-out examples.'),
      ('Interaction graph statistics','YES, descriptively','These count adjacent annotated speaker turns; they do not establish social relations.'),
      ('Leakage-safe evaluation','NOT YET','Keep evaluation gold separate from rule development/tuning; document sample selection and adjudication.'),
    ]: protocol.append(row)
    # style and widths
    for ws in wb.worksheets:
        ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
        for cell in ws[1]: cell.font=Font(bold=True,color='FFFFFF'); cell.fill=PatternFill('solid',fgColor='23415A'); cell.alignment=Alignment(wrap_text=True,vertical='top')
        for col in range(1,ws.max_column+1):
            vals=[len(str(ws.cell(r,col).value or '')) for r in range(1,min(ws.max_row,200)+1)]
            ws.column_dimensions[get_column_letter(col)].width=min(max(max(vals or [12])+2,12),55)
        for row in ws.iter_rows(min_row=2):
            for cell in row: cell.alignment=Alignment(vertical='top',wrap_text=True)
    wb.save(op)
    src.close()
    print('LITMIND AUTOMATED QA REPORT COMPLETE')
    print(f'Corpus rows analyzed: {len(data)}')
    print(f'Duplicate Speech_IDs: {len(duplicates)}; blank speakers: {len(blankspeaker)}; blank text: {len(blanktext)}; unmapped speaker labels: {len(unknown)}')
    print('No accuracy/precision/recall/F1 reported without independent gold labels and predictions.')
    print(f'Report: {op.resolve()}')
if __name__=='__main__': main()
