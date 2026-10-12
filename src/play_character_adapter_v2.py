"""
LitMind play-specific character adapter v2 (separate from frozen E3).
Uses the annotated CHARACTER_MASTER and UTTERANCE_CORPUS as the reference for speaker identity.
Does not import or modify nlp_processor.py or prev.py.

Run from D:\\LitMind:
  python src/play_character_adapter.py --corpus Dattani_Tara_Corpus_Final.xlsx --out reports/play_character_adapter
Requires: openpyxl, spacy, en_core_web_sm
"""
from pathlib import Path
from collections import Counter, defaultdict
import argparse
import re
import openpyxl
import spacy

PLAY_NAMES = {"D01": "Tara", "D02": "Dance Like a Man"}
STAGE_ACTION_WORDS = {
  "listens", "laughs", "smiles", "mutters", "yawns", "grins", "sighs",
  "weeps", "cries", "pauses", "moves", "turns", "looks", "sits", "stands"
}
# Conservative aliases: only map variants supported by the corpus's character names/notes
# or previously inspected NER surfaces. These are candidate mappings for review, not truth labels.
ALIASES = {
  "dan": "Dan",
  "chandan": "Chandan",
  "chandu": "Chandan",
  "tara": "Tara",
  "bharati": "Bharati",
  "mrs bharati patel": "Bharati",
  "patel": "Patel",
  "roopa": "Roopa",
  "dr thakkar": "Dr Thakkar",
  "thakkar": "Dr Thakkar",
  "jairaj": "Jairaj",
  "jai": "Jairaj",
  "jairaj parekh": "Jairaj",
  "ratna": "Ratna",
  "ratna devi": "Ratna",
  "viswas": "Viswas",
  "vishy": "Viswas",
  "lata": "Lata",
  "lata parekh": "Lata",
  "amritlal": "Amritlal",
  "amritlal parekh": "Amritlal",
}

def clean(s):
  return re.sub(r"\s+", " ", str(s or "")).strip()

def norm(s):
  s = clean(s).replace("’", "'").casefold()
  s = re.sub(r"[^a-zà-ÿ.' -]", "", s)
  return re.sub(r"\s+", " ", s).strip(" .'-")

def play_key(play_id, play_name=""):
  pid = clean(play_id).upper()
  if pid in PLAY_NAMES:
    return PLAY_NAMES[pid]
  text = (pid + " " + clean(play_name)).casefold()
  if "tara" in text:
    return "Tara"
  if "dance" in text or "man" in text:
    return "Dance Like a Man"
  return None


# Explicit variants are used for candidate discovery only. They are not relationship labels.
ALIAS_VARIANTS = {
  "Dan": ["Dan"],
  "Chandan": ["Chandan", "Chandu"],
  "Tara": ["Tara"],
  "Bharati": ["Bharati", "Mrs Bharati Patel"],
  "Patel": ["Patel", "Mr Patel"],
  "Roopa": ["Roopa", "Rupa"],
  "Dr Thakkar": ["Dr Thakkar", "Thakkar", "Doctor Thakkar"],
  "Jairaj": ["Jairaj", "Jai", "Jairaj Parekh"],
  "Ratna": ["Ratna", "Ratna Devi"],
  "Viswas": ["Viswas", "Vishy"],
  "Lata": ["Lata", "Lata Parekh"],
  "Amritlal": ["Amritlal", "Amritlal Parekh"],
}

def literal_alias_hits(text, canonical, variants, limit=3):
  """Find literal alias strings and example contexts; does not assert every hit is a character mention."""
  found = []
  total = 0
  for variant in variants:
    if not variant:
      continue
    pattern = r"(?<![A-Za-z])" + r"\s+".join(re.escape(part) for part in variant.split()) + r"(?![A-Za-z])"
    hits = list(re.finditer(pattern, text, flags=re.IGNORECASE))
    total += len(hits)
    for hit in hits[:max(0, limit-len(found))]:
      start = text.rfind("\n", 0, hit.start()) + 1
      end = text.find("\n", hit.end())
      if end < 0:
        end = min(len(text), hit.end()+180)
      context = re.sub(r"\s+", " ", text[start:end]).strip()
      if not context:
        context = re.sub(r"\s+", " ", text[max(0,hit.start()-100):min(len(text),hit.end()+100)]).strip()
      found.append((variant, context[:350]))
      if len(found) >= limit:
        break
  return total, found


def main():
  ap = argparse.ArgumentParser()
  ap.add_argument("--corpus", required=True)
  ap.add_argument("--out", default="reports/play_character_adapter")
  args = ap.parse_args()

  corpus = Path(args.corpus)
  out = Path(args.out)
  out.parent.mkdir(parents=True, exist_ok=True)
  if not corpus.exists():
    raise SystemExit(f"Corpus workbook not found: {corpus}")

  nlp = spacy.load("en_core_web_sm")
  source = openpyxl.load_workbook(corpus, read_only=True, data_only=True)
  if "CHARACTER_MASTER" not in source.sheetnames or "UTTERANCE_CORPUS" not in source.sheetnames:
    raise SystemExit("Workbook must contain CHARACTER_MASTER and UTTERANCE_CORPUS sheets.")

  # Read reference characters and their annotated speaker labels.
  master_ws = source["CHARACTER_MASTER"]
  master_rows = list(master_ws.iter_rows(values_only=True))
  master_headers = [clean(x) for x in master_rows[0]]
  required_master = ["Character_ID", "Play_ID", "Character_Name", "Source_Speaker_Label"]
  if any(x not in master_headers for x in required_master):
    raise SystemExit(f"Unexpected CHARACTER_MASTER headers: {master_headers}")
  mi = {name: master_headers.index(name) for name in required_master}
  characters = {}
  label_to_character = {}
  for row in master_rows[1:]:
    if not row:
      continue
    play = play_key(row[mi["Play_ID"]])
    name = clean(row[mi["Character_Name"]])
    label = clean(row[mi["Source_Speaker_Label"]]).upper()
    cid = clean(row[mi["Character_ID"]])
    if play and name:
      characters[(play, name)] = {"id": cid, "label": label}
      label_to_character[(play, label)] = name

  # Read annotated dialogue turns.
  utter_ws = source["UTTERANCE_CORPUS"]
  utter_rows = list(utter_ws.iter_rows(values_only=True))
  utter_headers = [clean(x) for x in utter_rows[0]]
  required_utter = ["Speech_ID", "Play_ID", "Play", "Speaker", "Speech_OCR"]
  if any(x not in utter_headers for x in required_utter):
    raise SystemExit(f"Unexpected UTTERANCE_CORPUS headers: {utter_headers}")
  ui = {name: utter_headers.index(name) for name in required_utter}

  turns = defaultdict(list)
  for row in utter_rows[1:]:
    if not row:
      continue
    play = play_key(row[ui["Play_ID"]], row[ui["Play"]])
    if play:
      turns[play].append({
        "speech_id": clean(row[ui["Speech_ID"]]),
        "speaker_label": clean(row[ui["Speaker"]]).upper(),
        "text": clean(row[ui["Speech_OCR"]]),
      })
  source.close()

  result = openpyxl.Workbook()
  summary = result.active
  summary.title = "Summary"
  summary.append(["Play", "Reference characters", "Utterance rows", "Rows with known speaker label",
                  "Unique annotated speakers", "PERSON mentions", "Unique PERSON surfaces",
                  "PERSON surfaces mapped to reference", "Distinct reference characters mapped",
                  "Action-like PERSON surfaces flagged"])
  speakers_ws = result.create_sheet("Speaker_Reference")
  speakers_ws.append(["Play", "Character_ID", "Reference character", "Source speaker label",
                      "Utterance rows with this speaker label"])
  ner_ws = result.create_sheet("NER_Candidates")
  ner_ws.append(["Play", "PERSON surface", "Mentions", "Mapped canonical character",
                 "Review flag", "Example context"])
  unmapped_ws = result.create_sheet("Unmapped_Speaker_Labels")
  unmapped_ws.append(["Play", "Speaker label in utterance corpus", "Utterance count", "Example Speech_ID"])
  alias_ws = result.create_sheet("Alias_Coverage")
  alias_ws.append(["Play", "Canonical reference character", "Candidate aliases searched",
                   "Literal text occurrences", "Example alias/context", "PERSON NER mentions mapped",
                   "Interpretation"])

  for play in ["Tara", "Dance Like a Man"]:
    ref_names = [name for (p, name) in characters if p == play]
    rows = turns[play]
    play_text = "\\n".join(row["text"] for row in rows)
    speaker_counts = Counter(row["speaker_label"] for row in rows if row["speaker_label"])
    known_rows = sum(1 for row in rows if (play, row["speaker_label"]) in label_to_character)
    for name in ref_names:
      info = characters[(play, name)]
      speakers_ws.append([play, info["id"], name, info["label"], speaker_counts.get(info["label"], 0)])

    for label, count in speaker_counts.items():
      if (play, label) not in label_to_character:
        example = next((r["speech_id"] for r in rows if r["speaker_label"] == label), "")
        unmapped_ws.append([play, label, count, example])

    surfaces = Counter()
    examples = {}
    # NER is used for name-mention candidates only; the annotated Speaker field is the
    # primary source for who is speaking in a given utterance.
    for row in rows:
      doc = nlp(row["text"])
      for ent in doc.ents:
        if ent.label_ != "PERSON":
          continue
        surface = clean(ent.text).strip(" .,'’")
        if not surface:
          continue
        surfaces[surface] += 1
        examples.setdefault(surface, ent.sent.text[:300])

    mapped_surfaces = {}
    for surface in surfaces:
      key = norm(surface)
      mapped = ALIASES.get(key)
      if not mapped:
        # Safe exact-name fallback, case-insensitive.
        for ref in ref_names:
          if key == norm(ref):
            mapped = ref
            break
      if mapped not in ref_names:
        mapped = None
      mapped_surfaces[surface] = mapped

    # Search canonical names/aliases directly in the corpus text, separately from spaCy NER.
    # This helps distinguish "name exists in text" from "NER recognized it as PERSON".
    for canonical in ref_names:
      variants = ALIAS_VARIANTS.get(canonical, [canonical])
      occurrence_count, occurrence_examples = literal_alias_hits(play_text, canonical, variants)
      ner_hits = [surface for surface, mapped_name in mapped_surfaces.items() if mapped_name == canonical]
      example_text = " || ".join(f"{variant}: {context}" for variant, context in occurrence_examples)
      if occurrence_count and not ner_hits:
        interpretation = "Alias/name occurs in text but no PERSON surface mapped; inspect contexts and NER miss."
      elif not occurrence_count and not ner_hits:
        interpretation = "No literal alias hit and no PERSON mapping; verify text coverage or add a verified variant."
      else:
        interpretation = "Text alias hits and/or NER mappings found; contexts still need manual review."
      alias_ws.append([play, canonical, "; ".join(variants), occurrence_count, example_text,
                       "; ".join(ner_hits), interpretation])

    distinct_mapped = {name for name in mapped_surfaces.values() if name}
    action_flag_count = 0
    for surface, count in surfaces.most_common():
      action_like = norm(surface) in STAGE_ACTION_WORDS
      if action_like:
        action_flag_count += 1
      flag = "REVIEW: action-like PERSON surface" if action_like else (
        "MAPPED alias/canonical candidate" if mapped_surfaces[surface] else "REVIEW: unmapped PERSON surface"
      )
      ner_ws.append([play, surface, count, mapped_surfaces[surface] or "", flag, examples.get(surface, "")])

    summary.append([
      play, len(ref_names), len(rows), known_rows, len(speaker_counts),
      sum(surfaces.values()), len(surfaces),
      sum(1 for value in mapped_surfaces.values() if value),
      len(distinct_mapped), action_flag_count
    ])

  xlsx = out.with_suffix(".xlsx")
  result.save(xlsx)
  print("\\nLITMIND PLAY CHARACTER ADAPTER COMPLETE")
  print("Frozen nlp_processor.py and prev.py were not imported or modified.")
  print(f"Report: {xlsx.resolve()}")
  for row in list(summary.values)[1:]:
    print(f"{row[0]}: reference characters={row[1]}, utterances={row[2]}, known-speaker rows={row[3]}, "
          f"PERSON mentions={row[5]}, unique PERSON surfaces={row[6]}, "
          f"mapped PERSON surfaces={row[7]}, distinct mapped characters={row[8]}")
  print("\\nImportant: Speaker_Reference uses the corpus's annotated Speaker column. "
        "NER_Candidates are mentions to review, not proof of relationships.")

if __name__ == "__main__":
  main()
