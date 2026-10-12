"""LitMind play-specific diagnostic (corrected alias matching + transparent review flags).
Does NOT import or modify nlp_processor.py.

Run from D:\\LitMind:
  python -u src/play_evaluation_fixed.py --corpus "Dattani_Tara_Corpus_Final.xlsx" --text-dir "data/processed" --out "reports/play_evaluation_fixed"
Requires: openpyxl, spacy, en_core_web_sm, networkx.
"""
from pathlib import Path
from collections import Counter
import argparse
import csv
import re
import openpyxl
import spacy
import networkx as nx  # retained as declared diagnostic dependency; not used to infer semantic relations

# Explicit aliases are reviewable assumptions, not claims that every occurrence is a character.
REFERENCE_ALIASES = {
  "Tara": {
    "Dan": ["Dan"],
    "Chandan": ["Chandan", "Chandu", "Chahdan", "Chahdan"],
    "Tara": ["Tara", "Tara Patel"],
    "Bharati": ["Bharati", "Bharati Patel", "Mrs Bharati Patel"],
    "Patel": ["Patel"],
    "Roopa": ["Roopa", "Rupa"],
    "Dr Thakkar": ["Dr Thakkar", "Thakkar"],
  },
  "Dance Like a Man": {
    "Jairaj": ["Jairaj", "Jai", "Jairaj Parekh", "Sri Jairaj Parekh"],
    "Ratna": ["Ratna", "Ratna Devi"],
    "Viswas": ["Viswas", "Vishwas", "Vishy"],
    "Lata": ["Lata", "Lata Parekh"],
    "Amritlal": ["Amritlal", "Amritlal Parekh"],
  },
}

# These are review flags only. They do not delete detections from the report.
ACTION_WORDS = {
  "listens", "laughs", "smiles", "grins", "mutters", "rubs", "follows",
  "realizes", "yawns", "drinks", "whispers", "surprised", "meaningfully",
  "quieter", "louder", "wears", "nudges", "hysterical", "funny",
}


def norm(s):
  s = str(s or "").replace("’", "'").replace("‘", "'").strip().lower()
  s = re.sub(r"\s+", " ", s)
  s = re.sub(r"[^a-zà-ÿ .'-]", "", s)
  return s.strip(" .'-")


def tokens(s):
  return [x for x in re.findall(r"[a-z]+", norm(s))
          if x not in {"mr", "mrs", "ms", "miss", "dr", "sir", "lady", "sri"}]


def play_key(play_id, play_name=""):
  pid = str(play_id or "").strip().upper()
  name = (str(play_id or "") + " " + str(play_name or "")).lower()
  if pid == "D01":
    return "Tara"
  if pid == "D02":
    return "Dance Like a Man"
  if "tara" in name:
    return "Tara"
  if "dance" in name or "man" in name:
    return "Dance Like a Man"
  return None


def match_reference(name, play):
  name_norm = norm(name)
  name_tokens = tokens(name)
  if not name_tokens:
    return None
  # First match exact canonical names / explicit aliases only.
  matches = []
  for canonical, aliases in REFERENCE_ALIASES[play].items():
    for alias in [canonical] + aliases:
      if name_norm == norm(alias):
        matches.append(canonical)
        break
  matches = sorted(set(matches))
  return matches[0] if len(matches) == 1 else None


def review_flag(name, context, matched):
  n = norm(name)
  c = (context or "").strip()
  if matched:
    return "Matched reference/alias; still review context"
  if n in ACTION_WORDS and (c.startswith("(") or f"({n}" in c.lower()):
    return "Likely stage-direction/action false positive"
  if c.startswith("(") and len(tokens(name)) == 1:
    return "Possible stage-direction false positive; review"
  if len(tokens(name)) == 1 and n in {"mummy", "dad", "grandfather", "uncle", "butler", "saar"}:
    return "Kinship/title/common noun; review"
  return "Unmatched PERSON surface; manual review"


def main():
  ap = argparse.ArgumentParser()
  ap.add_argument("--corpus", required=True, help="Path to Dattani corpus Excel workbook")
  ap.add_argument("--text-dir", required=True, help="Folder containing the two clean input .txt files")
  ap.add_argument("--out", default="reports/play_evaluation_fixed", help="Output prefix (without extension)")
  args = ap.parse_args()
  corpus, text_dir, out = Path(args.corpus), Path(args.text_dir), Path(args.out)
  out.parent.mkdir(parents=True, exist_ok=True)
  if not corpus.exists():
    raise SystemExit(f"Corpus workbook not found: {corpus}")
  if not text_dir.exists():
    raise SystemExit(f"Text directory not found: {text_dir}")

  nlp = spacy.load("en_core_web_sm")
  source_wb = openpyxl.load_workbook(corpus, read_only=True, data_only=True)
  required_sheets = {"CHARACTER_MASTER", "UTTERANCE_CORPUS"}
  missing_sheets = required_sheets - set(source_wb.sheetnames)
  if missing_sheets:
    raise SystemExit(f"Missing required workbook sheets: {', '.join(sorted(missing_sheets))}")

  master_rows = list(source_wb["CHARACTER_MASTER"].iter_rows(values_only=True))
  master_headers = [str(x or "").strip() for x in master_rows[0]]
  for key in ["Play_ID", "Character_Name"]:
    if key not in master_headers:
      raise SystemExit(f"CHARACTER_MASTER missing column {key!r}; found: {master_headers}")
  mi = {k: master_headers.index(k) for k in ["Play_ID", "Character_Name"]}
  refs = {"Tara": [], "Dance Like a Man": []}
  for row in master_rows[1:]:
    if not row:
      continue
    pk = play_key(row[mi["Play_ID"]])
    name = str(row[mi["Character_Name"]] or "").strip()
    if pk and name and name not in refs[pk]:
      refs[pk].append(name)

  utter_rows = list(source_wb["UTTERANCE_CORPUS"].iter_rows(values_only=True))
  utter_headers = [str(x or "").strip() for x in utter_rows[0]]
  needed = ["Play_ID", "Play", "Speech_ID", "Speaker", "Speech_OCR"]
  absent = [k for k in needed if k not in utter_headers]
  if absent:
    raise SystemExit(f"UTTERANCE_CORPUS missing columns {absent}; found: {utter_headers}")
  ui = {k: utter_headers.index(k) for k in needed}
  utterances = {"Tara": [], "Dance Like a Man": []}
  for row in utter_rows[1:]:
    if not row:
      continue
    pk = play_key(row[ui["Play_ID"]], row[ui["Play"]])
    if pk:
      utterances[pk].append((row[ui["Speech_ID"]], str(row[ui["Speaker"]] or "").strip(), str(row[ui["Speech_OCR"]] or "")))
  source_wb.close()

  workbook = openpyxl.Workbook()
  summary = workbook.active
  summary.title = "Summary"
  summary.append(["Play", "Reference speaker count", "Utterance rows", "Input text characters", "PERSON mentions", "Unique PERSON surfaces", "Unique surfaces matched to reference/alias", "Unique reference speakers matched", "Reference speakers not matched", "Unmatched PERSON surfaces"])
  det_ws = workbook.create_sheet("NER_Details")
  det_ws.append(["Play", "PERSON surface", "Mentions", "Matched reference speaker", "Review flag", "Example context"])
  miss_ws = workbook.create_sheet("Reference_Check")
  miss_ws.append(["Play", "Reference speaker", "NER match found?", "Matched surface(s)"])
  graph_ws = workbook.create_sheet("Interaction_Edges")
  graph_ws.append(["Play", "Speaker A", "Speaker B", "Adjacent-turn count", "Evidence"])

  totals = {}
  for play, filename in [("Tara", "Tara_clean_input.txt"), ("Dance Like a Man", "Dance_Like_a_Man_clean_input.txt")]:
    if len(refs[play]) == 0:
      raise SystemExit(f"No reference speakers loaded for {play}; check CHARACTER_MASTER Play_ID values.")
    path = text_dir / filename
    if not path.exists():
      raise SystemExit(f"Input text not found: {path}")
    text = path.read_text(encoding="utf-8")
    doc = nlp(text)
    surfaces, contexts = Counter(), {}
    for ent in doc.ents:
      if ent.label_ == "PERSON":
        name = re.sub(r"\s+", " ", ent.text.strip()).strip(" .,'’")
        if name:
          surfaces[name] += 1
          contexts.setdefault(name, ent.sent.text[:240])
    matched = {name: match_reference(name, play) for name in surfaces}
    seen_refs = {v for v in matched.values() if v}
    matched_surface_count = sum(1 for v in matched.values() if v)
    summary.append([play, len(refs[play]), len(utterances[play]), len(text), sum(surfaces.values()), len(surfaces), matched_surface_count, len(seen_refs), len(refs[play]) - len(seen_refs), sum(1 for v in matched.values() if not v)])
    for name, count in surfaces.most_common():
      det_ws.append([play, name, count, matched[name] or "", review_flag(name, contexts.get(name, ""), matched[name]), contexts.get(name, "")])
    for ref in refs[play]:
      hits = [name for name, val in matched.items() if val == ref]
      miss_ws.append([play, ref, "YES" if hits else "NO", "; ".join(hits)])

    rows = utterances[play]
    edge_counts, edge_evidence = Counter(), {}
    for i in range(1, len(rows)):
      _, prev_speaker, _ = rows[i - 1]
      speech_id, speaker, _ = rows[i]
      if not prev_speaker or not speaker or prev_speaker.casefold() == speaker.casefold():
        continue
      a, b = sorted([prev_speaker, speaker], key=str.casefold)
      edge_counts[(a, b)] += 1
      edge_evidence.setdefault((a, b), f"{rows[i-1][0]} → {speech_id}: {prev_speaker} / {speaker}")
    for (a, b), count in edge_counts.most_common():
      graph_ws.append([play, a, b, count, edge_evidence[(a, b)]])
    totals[play] = {"refs": refs[play], "utterances": len(rows), "mentions": sum(surfaces.values()), "unique": len(surfaces), "seen": seen_refs}

  xlsx = out.with_suffix(".xlsx")
  workbook.save(xlsx)
  csv_path = out.with_name(out.name + "_interaction_edges.csv")
  with csv_path.open("w", newline="", encoding="utf-8-sig") as f:
    writer = csv.writer(f)
    writer.writerow(["Play", "Speaker A", "Speaker B", "Adjacent-turn count", "Evidence"])
    for row in list(graph_ws.values)[1:]:
      writer.writerow(row)

  print("\nLITMIND PLAY DIAGNOSTIC COMPLETE (ALIAS-AWARE)")
  print("Original nlp_processor.py and prev.py were not imported or modified.")
  print(f"Report: {xlsx.resolve()}")
  print(f"Interaction edges: {csv_path.resolve()}")
  for play, data in totals.items():
    print(f"{play}: reference speakers={len(data['refs'])}, utterance rows={data['utterances']}, PERSON mentions={data['mentions']}, unique PERSON surfaces={data['unique']}, reference speakers matched={len(data['seen'])}/{len(data['refs'])}")
  print("\nReview flags are suggestions, not automatic removals. Adjacent-turn edges do NOT prove semantic relationships.")

if __name__ == "__main__":
  main()
